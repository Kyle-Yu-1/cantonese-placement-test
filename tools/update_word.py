#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把一次语音练习的评分结果追加进 Word 评分记录。
用法: python tools/update_word.py --payload <json路径>
payload: {attempt, ts, duration, transcript, scores[6], total, pct, level, flags[], missing[], weak, lever, docx?}
"""
import argparse, json, math, os, sys, tempfile, datetime

DEFAULT_DOCX = r"C:\Users\19461\OneDrive\Desktop\机械产品介绍AI学习评分记录.docx"
FONT_CANDIDATES = [r"C:\Windows\Fonts\msyh.ttc", r"C:\Windows\Fonts\simhei.ttf"]
DIMS = ["流利度", "准确性", "语言复杂度", "信息完整性", "逻辑结构", "自信度"]

def load_font(size):
    from PIL import ImageFont
    for f in FONT_CANDIDATES:
        if os.path.exists(f):
            try:
                return ImageFont.truetype(f, size)
            except Exception:
                pass
    return ImageFont.load_default()

def make_radar(path, scores):
    from PIL import Image, ImageDraw
    W = H = 800
    img = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(img)
    f_l = load_font(20); f_s = load_font(16)
    cx, cy, Rmax, n = 400, 410, 250, 6
    angles = [-90 + i * 60 for i in range(n)]
    def pt(i, v):
        r = (v - 1) / 4 * Rmax
        a = math.radians(angles[i])
        return (cx + r * math.cos(a), cy + r * math.sin(a))
    for v in range(1, 6):
        poly = [pt(i, v) for i in range(n)]
        if v > 1:
            d.line(poly + [poly[0]], fill=(210, 210, 210), width=1)
        x, y = pt(0, v)
        d.text((x - 26, y - 12), str(v), font=f_s, fill=(150, 150, 150))
    for i in range(n):
        d.line([(cx, cy), pt(i, 5)], fill=(215, 215, 215), width=1)
    for i, name in enumerate(DIMS):
        x, y = pt(i, 5)
        dx, dy = math.cos(math.radians(angles[i])), math.sin(math.radians(angles[i]))
        d.text((x + dx * 58, y + dy * 40 - 12), name, font=f_l, fill=(30, 30, 30), anchor="mm")
    poly = [pt(i, v) for i, v in enumerate(scores)]
    overlay = Image.new("RGBA", (W, H), (0, 0, 0, 0))
    od = ImageDraw.Draw(overlay)
    od.polygon(poly, fill=(47, 107, 178, 70))
    img = Image.alpha_composite(img.convert("RGBA"), overlay).convert("RGB")
    d = ImageDraw.Draw(img)
    d.line(poly + [poly[0]], fill=(47, 107, 178), width=4)
    for i, v in enumerate(scores):
        x, y = pt(i, v)
        d.ellipse([x - 6, y - 6, x + 6, y + 6], fill=(47, 107, 178))
        dx, dy = math.cos(math.radians(angles[i])), math.sin(math.radians(angles[i]))
        d.text((x + dx * 26, y + dy * 26 - 12), str(v), font=f_s, fill=(47, 107, 178), anchor="mm")
    img.save(path)
    return path

def set_font(run, size=None, bold=None, name="微软雅黑"):
    from docx.shared import Pt
    from docx.oxml.ns import qn
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rf = rPr.get_or_add_rFonts()
    rf.set(qn("w:ascii"), name); rf.set(qn("w:hAnsi"), name); rf.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold

def para(doc, text, size=11, bold=False, align=None, space_after=6):
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    p = doc.add_paragraph()
    r = p.add_run(text); set_font(r, size=size, bold=bold)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p

def heading(doc, text, level=1):
    from docx.shared import Pt
    p = doc.add_paragraph()
    r = p.add_run(text); set_font(r, size=15 if level == 1 else 12, bold=True)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(6)
    return p

def count_attempts(doc):
    n = 0
    for p in doc.paragraphs:
        if "次语音练习" in p.text:
            n += 1
    return n

def append_record(payload):
    from docx import Document
    from docx.shared import Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    docx_path = payload.get("docx") or DEFAULT_DOCX
    if not os.path.exists(docx_path):
        return {"ok": False, "error": "找不到 Word 文件：" + docx_path}
    try:
        doc = Document(docx_path)
    except Exception as e:
        return {"ok": False, "error": "无法打开 Word：" + str(e)}

    scores = [int(x) for x in (payload.get("scores") or [3] * 6)][:6]
    while len(scores) < 6:
        scores.append(3)
    total = int(payload.get("total") or sum(scores))
    pct = round(total / 30 * 1000) / 10
    level = payload.get("level") or ("卓越" if pct >= 90 else "熟练" if pct >= 80 else "达标" if pct >= 60 else "发展中" if pct >= 40 else "起步")

    n = int(payload.get("attempt") or 0)
    if n <= 0:
        n = count_attempts(doc) + 1
    try:
        ts = datetime.datetime.fromisoformat(payload.get("ts", "").replace("Z", "+00:00"))
        date_str = ts.astimezone().strftime("%Y-%m-%d %H:%M")
    except Exception:
        date_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")

    label = payload.get("label")
    heading(doc, label or ("第 %d 次语音练习 · %s" % (n, date_str)), 1)
    para(doc, "时长 %d 秒 · 总分 %d/30 · 百分制 %.1f · 等级 %s" % (int(payload.get("duration") or 0), total, pct, level), size=10, bold=True)
    para(doc, "转写文本：", size=10, bold=True, space_after=2)
    para(doc, str(payload.get("transcript") or "（无）"), size=10, space_after=8)

    dims = payload.get("dims")
    has_dims = isinstance(dims, list) and len(dims) >= 6
    if has_dims:
        rows = [("评价维度", "得分", "评语")] + [(DIMS[i], str(scores[i]), str(dims[i].get("note", "")) if isinstance(dims[i], dict) else str(dims[i])) for i in range(6)]
        rows += [("合计（满分 30）", str(total), ""), ("百分制得分", "%.1f" % pct, ""), ("等级判定", level, "")]
        table = doc.add_table(rows=0, cols=3)
        table.style = "Table Grid"
        widths = [Inches(1.4), Inches(0.6), Inches(4.5)]
        for r_vals in rows:
            cells = table.add_row().cells
            for ci, val in enumerate(r_vals):
                cell = cells[ci]
                cell.width = widths[ci]
                p = cell.paragraphs[0]
                run = p.add_run(val)
                set_font(run, size=10, bold=(r_vals[0] in ("评价维度", "合计（满分 30）", "百分制得分", "等级判定")))
                if ci == 1:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if ci == 2:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    else:
        rows = [("评价维度", "得分")] + [(DIMS[i], str(scores[i])) for i in range(6)]
        rows += [("合计（满分 30）", str(total)), ("百分制得分", "%.1f" % pct), ("等级判定", level)]
        table = doc.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        widths = [Inches(2.4), Inches(0.9)]
        for r_vals in rows:
            cells = table.add_row().cells
            for ci, val in enumerate(r_vals):
                cell = cells[ci]
                cell.width = widths[ci]
                p = cell.paragraphs[0]
                run = p.add_run(val)
                set_font(run, size=10, bold=(r_vals[0] in ("评价维度", "合计（满分 30）", "百分制得分", "等级判定")))
                if ci == 1:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    tmp_png = os.path.join(tempfile.gettempdir(), "pitch_radar_%d.png" % os.getpid())
    try:
        make_radar(tmp_png, scores)
        doc.add_picture(tmp_png, width=Inches(3.0))
        last = doc.paragraphs[-1]
        last.alignment = WD_ALIGN_PARAGRAPH.CENTER
    except Exception:
        pass

    flags = payload.get("flags") or []
    if flags:
        para(doc, "需纠正：", size=10, bold=True, space_after=2)
        for f in flags:
            p = doc.add_paragraph(style="List Bullet")
            r = p.add_run(str(f)); set_font(r, size=10)
    if payload.get("weak"):
        para(doc, "阿基米德杠杆问题：%s" % payload.get("lever", ""), size=10)

    para(doc, "（本记录由「机械产品介绍·语音训练」系统生成）", size=9)
    try:
        doc.save(docx_path)
    except PermissionError:
        return {"ok": False, "error": "Word 文件被占用，请先关闭 WPS/Word 后再试"}
    except Exception as e:
        return {"ok": False, "error": "保存失败：" + str(e)}
    return {"ok": True, "path": docx_path, "attempt": n, "total": total, "pct": pct, "level": level}

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--payload", required=True)
    args = ap.parse_args()
    try:
        with open(args.payload, "r", encoding="utf-8") as f:
            payload = json.load(f)
    except Exception as e:
        print(json.dumps({"ok": False, "error": "读取请求失败：" + str(e)}, ensure_ascii=False))
        return
    result = append_record(payload)
    print(json.dumps(result, ensure_ascii=False))

if __name__ == "__main__":
    main()
